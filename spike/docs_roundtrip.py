"""Round-trip and safety tests for service/docs.py. Run: .venv/Scripts/python spike/docs_roundtrip.py"""

import copy
import io
import sys
import zipfile

sys.path.insert(0, "service")
import docs  # noqa: E402
from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

W = docs.W
OUT = "spike/out/docs_mine/"
checks = 0


def ok(cond, msg):
    global checks
    assert cond, msg
    checks += 1


def resolve(docx_bytes: bytes, accept: bool) -> list[str]:
    """Paragraph texts after Word's Accept All (accept=True) or Reject All (accept=False)."""
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    root = etree.fromstring(z.read("word/document.xml"))
    root = copy.deepcopy(root)
    drop, keep = (W + "del", W + "ins") if accept else (W + "ins", W + "del")
    for el in list(root.iter(drop)):
        el.getparent().remove(el)
    for el in list(root.iter(keep)):
        parent = el.getparent()
        i = parent.index(el)
        for ch in list(el):
            parent.insert(i, ch)
            i += 1
        parent.remove(el)
    for t in root.iter(W + "delText"):
        t.tag = W + "t"
    paras = []
    for p in root.iter(W + "p"):
        paras.append("".join(t.text or "" for t in p.iter(W + "t")))
    return paras


# ---- PPTX: reference split across runs, a table, notes, change log in notes
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[5])
p = s.shapes.title.text_frame.paragraphs[0]
r1 = p.add_run(); r1.text = "Summer starts at Gre"
r2 = p.add_run(); r2.text = "ggs"; r2.font.bold = True
r3 = p.add_run(); r3.text = " with friends."; r3.font.italic = True
s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.text = "Play Coronation Street on the big screen."
tbl = s.shapes.add_table(2, 2, Inches(1), Inches(3.5), Inches(6), Inches(1)).table
tbl.cell(0, 0).text = "Snack"; tbl.cell(0, 1).text = "Nando's peri-peri"
s.notes_slide.notes_text_frame.text = "Mention Tim Hortons in the toast."
s2 = prs.slides.add_slide(prs.slide_layouts[1]); s2.shapes.title.text = "Slide two"; s2.placeholders[1].text = "Wawa runs at midnight."
b = io.BytesIO(); prs.save(b); pptx_bytes = b.getvalue(); open(OUT + "in.pptx", "wb").write(pptx_bytes)
parsed = docs.parse("deck.pptx", pptx_bytes)
changes = []
for blk in parsed["blocks"]:
    t = blk["text"]
    if "Greggs" in t: changes.append({"id": blk["id"], "text": t.replace("Greggs", "Chick-fil-A"), "why": "Atlanta's favourite"})
    if "Nando's" in t: changes.append({"id": blk["id"], "text": t.replace("Nando's", "Zaxby's")})
    if "Tim Hortons" in t: changes.append({"id": blk["id"], "text": "Raise a glass in the toast."})
name, out, media = docs.export("deck.pptx", pptx_bytes, changes, audience="Atlanta")
open(OUT + name, "wb").write(out)
after = [x["text"] for x in docs.parse(name, out)["blocks"]]
prs2 = Presentation(io.BytesIO(out))
runs = [(r.text, r.font.bold, r.font.italic) for r in prs2.slides[0].shapes.title.text_frame.paragraphs[0].runs]
notes = prs2.slides[0].notes_slide.notes_text_frame.text
ok(any("Chick-fil-A" in x for x in after) and any("Zaxby" in x for x in after), "pptx swaps land")
ok(runs[-1] == (" with friends.", None, True), "italic tail run survives")
ok("Raise a glass" in notes, "notes paragraph rewritten")
ok("Room · Atlanta changed this slide:" in notes and "Atlanta's favourite" in notes, f"change log in notes: {notes!r}")
ok("changed this slide" not in prs2.slides[1].notes_slide.notes_text_frame.text if prs2.slides[1].has_notes_slide else True, "untouched slide gets no log")

# ---- DOCX: tracked changes with a reason; Accept All and Reject All both exact
d = Document()
para = d.add_paragraph(); para.add_run("We met at "); rr = para.add_run("Wetherspoons"); rr.bold = True; para.add_run(" in Peckham.")
d.add_paragraph("He still quotes Love Actually every Christmas.")
d.add_paragraph("Tonight we raise a Guinness.")
t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "Team"; t.cell(0, 1).text = "Arsenal forever"
b = io.BytesIO(); d.save(b); docx_bytes = b.getvalue()
parsed = docs.parse("toast.docx", docx_bytes)
original = [x["text"] for x in parsed["blocks"]]
want = {
    "p0": ("We met at a Nando's in Peckham.", "A chain the room knows"),
    "p2": ("Tonight we raise a glass.", "Cut for a Lagos crowd"),
    "t0.0.1.0": ("Gor Mahia forever", ""),
}
changes = [{"id": k, "text": v[0], "why": v[1]} for k, v in want.items()]
name, out, media = docs.export("toast.docx", docx_bytes, changes, audience="Lagos")
open(OUT + "toast (Room).docx", "wb").write(out)
acc, rej = resolve(out, True), resolve(out, False)
ok("We met at a Nando's in Peckham." in acc and "Tonight we raise a glass." in acc and "Gor Mahia forever" in acc, f"accept all: {acc}")
ok("We met at Wetherspoons in Peckham." in rej and "Tonight we raise a Guinness." in rej and "Arsenal forever" in rej, f"reject all: {rej}")
z = zipfile.ZipFile(io.BytesIO(out))
ok("word/comments.xml" in z.namelist(), "comments part written")
cx = z.read("word/comments.xml").decode()
ok("A chain the room knows" in cx and "Cut for a Lagos crowd" in cx, "both reasons present")
dx = z.read("word/document.xml").decode()
ok('w:author="Room · Lagos"' in dx, "revisions carry the author")
d2 = Document(io.BytesIO(out))
bold_ins = [r for r in d2.paragraphs[0]._p.iter(W + "ins") for r in r.iter(W + "r") if r.find(f"{W}rPr/{W}b") is not None]
ok(bool(bold_ins), "inserted text keeps the bold of the brand it replaced")

# ---- PDF to tracked DOCX
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    st = getSampleStyleSheet()
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build([Paragraph("Launch brief", st["Title"]),
                                               Paragraph("Close with Tim Hortons coffee for the crew.", st["Normal"])])
    pdf = buf.getvalue()
except ImportError:
    pdf = open(OUT + "brief.pdf", "rb").read()
parsed = docs.parse("brief.pdf", pdf)
ch = [{"id": x["id"], "text": x["text"].replace("Tim Hortons", "Chick-fil-A"), "why": "Local"} for x in parsed["blocks"] if "Tim Hortons" in x["text"]]
name, out, _ = docs.export("brief.pdf", pdf, ch)
ok(name.endswith(".docx"), "pdf comes back as docx")
ok(any("Chick-fil-A coffee" in x for x in resolve(out, True)) and any("Tim Hortons coffee" in x for x in resolve(out, False)), "pdf edits tracked")

# ---- TXT keeps spacing
txt = "First line with Greggs.\n\n\nSecond paragraph.\n"
name, out, _ = docs.export("set.txt", txt.encode(), [{"id": "b0", "text": "First line with Chuy's."}])
ok(out.decode() == "First line with Chuy's.\n\n\nSecond paragraph.\n", "txt spacing kept")

# ---- hostile files
def refused(name, data, needle):
    try:
        docs.parse(name, data)
    except docs.DocError as e:
        ok(needle in str(e), f"{name}: {e}")
        return
    raise AssertionError(f"{name} was accepted")

bomb = io.BytesIO()
with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
    zf.writestr("[Content_Types].xml", "<x/>")
    zf.writestr("word/document.xml", "A" * (40 * 1024 * 1024))
refused("bomb.docx", bomb.getvalue(), "too large")
macro = io.BytesIO()
with zipfile.ZipFile(macro, "w") as zf:
    zf.writestr("[Content_Types].xml", "<x/>")
    zf.writestr("word/vbaProject.bin", b"\x00" * 10)
refused("macro.docx", macro.getvalue(), "macros")
refused("fake.pptx", b"%PDF-1.4 not a deck", "damaged")
refused("empty.txt", b"   \n\n  ", "no text")
try:
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(200, 200); w.encrypt("secret")
    eb = io.BytesIO(); w.write(eb)
    refused("locked.pdf", eb.getvalue(), "password-locked")
except ImportError:
    pass
try:
    docs.parse("evil.exe", b"MZ")
    raise AssertionError("exe accepted")
except docs.DocError:
    checks += 1

print(f"ALL OK ({checks} checks)")

# ---- a Room-edited Word file reads back with its tracked insertions, and can be edited again
name, out, _ = docs.export("toast.docx", docx_bytes, [{"id": "p0", "text": "We met at a Nando's in Peckham.", "why": "x"}])
again = docs.parse(name, out)
ok(again["blocks"][0]["text"] == "We met at a Nando's in Peckham.", f"re-read sees insertions: {again['blocks'][0]}")
name2, out2, _ = docs.export(name, out, [{"id": "p0", "text": "We met at a Chicken Republic in Peckham.", "why": "y"}])
ok("We met at a Chicken Republic in Peckham." in resolve(out2, True), f"second pass: {resolve(out2, True)}")
print(f"ALL OK after re-edit ({checks} checks)")
